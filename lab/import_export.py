import csv
import io
import os
import uuid
from pathlib import Path
from typing import Any

from django.conf import settings
from django.db import transaction
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import Student, Class, Enrollment, Device, Classroom, MediaFolder, MediaAsset

# Supported import entities
IMPORT_ENTITIES = [
    ("students", _("Students")),
    ("classes", _("Classes")),
    ("enrollments", _("Enrollments")),
]

DUPLICATE_MODES = [
    ("reject", _("Reject duplicates")),
    ("update", _("Update existing")),
    ("skip", _("Skip duplicates")),
]

# Required fields per entity
REQUIRED_FIELDS = {
    "students": ["student_id", "full_name"],
    "classes": ["code", "name", "grade"],
    "enrollments": ["student_id", "class_code"],
}

# Allowed fields per entity
ALLOWED_FIELDS = {
    "students": ["student_id", "full_name", "active", "preferred_language", "device_number", "classroom_code"],
    "classes": ["code", "name", "grade", "section", "learning_language", "academic_year", "active"],
    "enrollments": ["student_id", "class_code", "status"],
}


class ParseError(Exception):
    pass


def _safe_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _to_bool(value: Any) -> bool | None:
    if value is None:
        return None
    val = _safe_str(value).lower()
    if val in ("true", "1", "yes", "y", "نعم", "oui"):
        return True
    if val in ("false", "0", "no", "n", "لا", "non"):
        return False
    return None


def _normalize_status(value: Any) -> str | None:
    val = _safe_str(value).lower()
    mapping = {
        "active": "active",
        "نشط": "active",
        "actif": "active",
        "inactive": "inactive",
        "غير نشط": "inactive",
        "inactif": "inactive",
        "completed": "completed",
        "منجز": "completed",
        "terminé": "completed",
        "dropped": "dropped",
        "منسحب": "dropped",
        "abandonné": "dropped",
    }
    return mapping.get(val, val if val in dict(Enrollment.status.field.choices) else None)


def parse_uploaded_file(uploaded_file) -> tuple[str, list[dict]]:
    """Parse uploaded CSV or XLSX file. Returns (entity_type, rows)."""
    name = uploaded_file.name.lower()
    content = uploaded_file.read()
    if not content:
        raise ParseError(_("Empty file."))

    if name.endswith(".csv"):
        return _parse_csv(content)
    elif name.endswith((".xlsx", ".xls")):
        return _parse_xlsx(content)
    else:
        raise ParseError(_("Unsupported file format. Use CSV or XLSX."))


def _detect_entity_type(headers: list[str]) -> str:
    normalized = [h.lower().strip() for h in headers]
    for entity, fields in REQUIRED_FIELDS.items():
        if all(f in normalized for f in fields):
            return entity
    raise ParseError(_("Could not detect entity type. Required fields missing: %(fields)s.") % {
        "fields": ", ".join(["student_id, full_name", "code, name, grade", "student_id, class_code"])
    })


def _parse_csv(content: bytes) -> tuple[str, list[dict]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ParseError(_("File is not valid UTF-8."))

    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise ParseError(_("Empty CSV or missing header row."))

    headers = [h.strip() for h in reader.fieldnames if h]
    entity_type = _detect_entity_type(headers)

    rows = []
    for row_num, row in enumerate(reader, start=2):
        normalized = {k.strip(): _safe_str(v) for k, v in row.items() if k}
        rows.append(normalized)
    return entity_type, rows


def _parse_xlsx(content: bytes) -> tuple[str, list[dict]]:
    try:
        import zipfile
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise ParseError(_("Invalid XLSX file."))
    except Exception:
        raise ParseError(_("Invalid XLSX file."))

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise ParseError(_("Empty XLSX file."))

    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(cell).strip() if cell is not None else "" for cell in next(rows_iter)]
    except StopIteration:
        raise ParseError(_("Empty XLSX or missing header row."))

    headers = [h for h in headers if h]
    entity_type = _detect_entity_type(headers)

    rows = []
    for row in rows_iter:
        data = {headers[i]: _safe_str(row[i]) if i < len(row) else "" for i in range(len(headers))}
        rows.append(data)
    return entity_type, rows


def validate_rows(entity_type: str, rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Validate rows. Returns (valid_rows, error_rows)."""
    valid_rows = []
    error_rows = []

    required = REQUIRED_FIELDS.get(entity_type, [])
    allowed = ALLOWED_FIELDS.get(entity_type, [])

    for idx, row in enumerate(rows, start=2):
        errors = []
        normalized = {}

        for field in allowed:
            if field in row:
                normalized[field] = _safe_str(row[field])

        for req_field in required:
            if not normalized.get(req_field):
                errors.append(_("Missing required field: %(field)s") % {"field": req_field})

        if errors:
            error_rows.append({"row": idx, "data": row, "errors": errors})
            continue

        valid_rows.append(normalized)

    return valid_rows, error_rows


def _validate_student_row(row: dict) -> list[str]:
    errors = []
    if Student.objects.filter(student_id=row["student_id"]).exists():
        errors.append(_("Duplicate student ID: %(id)s") % {"id": row["student_id"]})
    if row.get("device_number"):
        try:
            device_num = int(row["device_number"])
            if device_num < 1 or device_num > 40:
                errors.append(_("Device number must be 1-40."))
        except ValueError:
            errors.append(_("Invalid device number."))
    if row.get("classroom_code"):
        if not Classroom.objects.filter(code=row["classroom_code"], active=True).exists():
            errors.append(_("Invalid classroom code: %(code)s") % {"code": row["classroom_code"]})
    if row.get("preferred_language") and row["preferred_language"] not in ["en", "ar", "fr"]:
        errors.append(_("Invalid language code: %(lang)s") % {"lang": row["preferred_language"]})
    return errors


def _validate_class_row(row: dict) -> list[str]:
    errors = []
    if Class.objects.filter(code=row["code"]).exists():
        errors.append(_("Duplicate class code: %(code)s") % {"code": row["code"]})
    if row.get("learning_language") and row["learning_language"] not in ["en", "ar", "fr"]:
        errors.append(_("Invalid language code: %(lang)s") % {"lang": row["learning_language"]})
    return errors


def _validate_enrollment_row(row: dict) -> list[str]:
    errors = []
    if not Student.objects.filter(student_id=row["student_id"]).exists():
        errors.append(_("Student not found: %(id)s") % {"id": row["student_id"]})
    if not Class.objects.filter(code=row["class_code"]).exists():
        errors.append(_("Class not found: %(code)s") % {"code": row["class_code"]})
    if Enrollment.objects.filter(student__student_id=row["student_id"], class_obj__code=row["class_code"]).exists():
        errors.append(_("Duplicate enrollment: %(student)s in %(class)s") % {
            "student": row["student_id"], "class": row["class_code"]
        })
    if row.get("status"):
        valid_statuses = [s[0] for s in Enrollment.status.field.choices]
        if row["status"] not in valid_statuses:
            errors.append(_("Invalid enrollment status: %(status)s") % {"status": row["status"]})
    return errors


def deep_validate(entity_type: str, valid_rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Deep validation against database. Returns (clean_rows, error_rows)."""
    clean_rows = []
    error_rows = []

    validators = {
        "students": _validate_student_row,
        "classes": _validate_class_row,
        "enrollments": _validate_enrollment_row,
    }
    validator = validators.get(entity_type)

    for idx, row in enumerate(valid_rows, start=2):
        if validator:
            errors = validator(row)
        else:
            errors = []

        if errors:
            error_rows.append({"row": idx, "data": row, "errors": errors})
        else:
            clean_rows.append(row)

    return clean_rows, error_rows


def perform_import(entity_type: str, clean_rows: list[dict], duplicate_mode: str, user) -> dict:
    """Perform transactional import. Returns report dict."""
    report = {
        "total": len(clean_rows),
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "errors": [],
    }

    with transaction.atomic():
        for row in clean_rows:
            try:
                if entity_type == "students":
                    _import_student(row, duplicate_mode, report, user)
                elif entity_type == "classes":
                    _import_class(row, duplicate_mode, report, user)
                elif entity_type == "enrollments":
                    _import_enrollment(row, duplicate_mode, report, user)
            except Exception as e:
                report["errors"].append({"row": row, "error": str(e)})
                raise  # Rollback transaction

    return report


def _import_student(row: dict, duplicate_mode: str, report: dict, user):
    student_id = row["student_id"]
    defaults = {
        "full_name": row.get("full_name", student_id),
        "active": _to_bool(row.get("active")) if row.get("active") else True,
        "preferred_language": row.get("preferred_language", ""),
    }
    if row.get("device_number"):
        try:
            device = Device.objects.get(number=int(row["device_number"]), classroom__code=row.get("classroom_code", ""))
            defaults["assigned_device"] = device
        except Device.DoesNotExist:
            pass

    existing = Student.objects.filter(student_id=student_id).first()
    if existing:
        if duplicate_mode == "update":
            for key, value in defaults.items():
                setattr(existing, key, value)
            existing.save()
            report["updated"] += 1
        elif duplicate_mode == "skip":
            report["skipped"] += 1
        else:  # reject
            report["skipped"] += 1
    else:
        Student.objects.create(student_id=student_id, **defaults)
        report["created"] += 1


def _import_class(row: dict, duplicate_mode: str, report: dict, user):
    code = row["code"]
    defaults = {
        "name": row.get("name", code),
        "grade": row.get("grade", ""),
        "section": row.get("section", ""),
        "learning_language": row.get("learning_language", ""),
        "academic_year": row.get("academic_year", ""),
        "active": _to_bool(row.get("active")) if row.get("active") else True,
    }
    existing = Class.objects.filter(code=code).first()
    if existing:
        if duplicate_mode == "update":
            for key, value in defaults.items():
                setattr(existing, key, value)
            existing.save()
            report["updated"] += 1
        elif duplicate_mode == "skip":
            report["skipped"] += 1
        else:
            report["skipped"] += 1
    else:
        Class.objects.create(code=code, **defaults)
        report["created"] += 1


def _import_enrollment(row: dict, duplicate_mode: str, report: dict, user):
    student_id = row["student_id"]
    class_code = row["class_code"]
    status = _normalize_status(row.get("status")) or "active"

    try:
        student = Student.objects.get(student_id=student_id)
        class_obj = Class.objects.get(code=class_code)
    except (Student.DoesNotExist, Class.DoesNotExist):
        report["errors"].append({"row": row, "error": _("Student or class not found.")})
        return

    existing = Enrollment.objects.filter(student=student, class_obj=class_obj).first()
    if existing:
        if duplicate_mode == "update":
            existing.status = status
            existing.save()
            report["updated"] += 1
        elif duplicate_mode == "skip":
            report["skipped"] += 1
        else:
            report["skipped"] += 1
    else:
        Enrollment.objects.create(student=student, class_obj=class_obj, status=status)
        report["created"] += 1


def export_csv(entity_type: str, queryset) -> bytes:
    """Export data as CSV bytes."""
    buffer = io.StringIO()
    if entity_type == "students":
        writer = csv.writer(buffer)
        writer.writerow(["student_id", "full_name", "active", "preferred_language", "device_number", "classroom_code"])
        for s in queryset.select_related("assigned_device", "assigned_device__classroom"):
            writer.writerow([
                s.student_id,
                s.full_name,
                "true" if s.active else "false",
                s.preferred_language,
                s.assigned_device.number if s.assigned_device else "",
                s.assigned_device.classroom.code if s.assigned_device and s.assigned_device.classroom else "",
            ])
    elif entity_type == "classes":
        writer = csv.writer(buffer)
        writer.writerow(["code", "name", "grade", "section", "learning_language", "academic_year", "active"])
        for c in queryset:
            writer.writerow([
                c.code, c.name, c.grade, c.section, c.learning_language, c.academic_year,
                "true" if c.active else "false",
            ])
    elif entity_type == "enrollments":
        writer = csv.writer(buffer)
        writer.writerow(["student_id", "class_code", "status"])
        for e in queryset.select_related("student", "class_obj"):
            writer.writerow([e.student.student_id, e.class_obj.code, e.status])
    return buffer.getvalue().encode("utf-8-sig")


def export_xlsx(entity_type: str, queryset) -> bytes:
    """Export data as XLSX bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = entity_type.capitalize()

    headers = {
        "students": ["student_id", "full_name", "active", "preferred_language", "device_number", "classroom_code"],
        "classes": ["code", "name", "grade", "section", "learning_language", "academic_year", "active"],
        "enrollments": ["student_id", "class_code", "status"],
    }
    ws.append(headers.get(entity_type, []))

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    if entity_type == "students":
        for s in queryset.select_related("assigned_device", "assigned_device__classroom"):
            ws.append([
                s.student_id,
                s.full_name,
                "Yes" if s.active else "No",
                s.preferred_language,
                s.assigned_device.number if s.assigned_device else "",
                s.assigned_device.classroom.code if s.assigned_device and s.assigned_device.classroom else "",
            ])
    elif entity_type == "classes":
        for c in queryset:
            ws.append([
                c.code, c.name, c.grade, c.section, c.learning_language, c.academic_year,
                "Yes" if c.active else "No",
            ])
    elif entity_type == "enrollments":
        for e in queryset.select_related("student", "class_obj"):
            ws.append([e.student.student_id, e.class_obj.code, e.status])

    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
